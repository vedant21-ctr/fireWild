"""
build_deck.py - Render the Phase 1 slide deck to PDF and PPTX.

The slides in presentation/slides/*.html are the source of the online deck
(https://claude.ai/artifact/TAmZVEaNQ152R9mQyRN816). This script assembles them in
deck.json order, points the figures at experiments/ndws/, and uses headless Google
Chrome to write:
    presentation/fireWild_phase1.pdf    vector PDF, one 16:9 page per slide
    presentation/fireWild_phase1.pptx   16:9 PowerPoint, one full-slide image per slide,
                                        speaker notes included

Usage (macOS, Google Chrome installed):
    python presentation/build_deck.py
"""

import os
import re
import json
import shutil
import subprocess
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(HERE)
CHROME = os.environ.get("CHROME", "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
W, H = 1920, 1080

# uploaded-asset ids used by the online deck -> repository figures
FIGURES = {
    "a6597fc4e5b1a9a51ccce4071dfaabaf": "experiments/ndws/eda/example_patches.png",
    "087d78f1e2f122d435f0df55b5dcf199": "experiments/ndws/eda/spread_vs_distance.png",
    "60329f20556fe30d306e75ce4beee656": "experiments/ndws/eda/wind_orientation.png",
    "82f2f8b4679efd6bd89fbe6e2db9fc6d": "experiments/ndws/eda/split_overlap.png",
    "b061f6fcb64e66cf2d5116b4a0275070": "experiments/ndws/eda/channel_distributions.png",
    "1504ae2789086ec7c519c63e1f42b157": "experiments/ndws/results/protocol_gap.png",
    "d78860bbbcd8e8b01e6aa0fa0f835378": "experiments/ndws/results/spatial/pr_curves.png",
    "1ac0d87dea927f5d74550b2080f15248": "experiments/ndws/results/spatial/probability_maps.png",
    "a14ec5f5477cf5380109b95f38b13fb3": "experiments/ndws/results/spatial/shap_summary.png",
    "b9321dac522b831bd207f2160a0870e6": "experiments/ndws/results/spatial/plausibility.png",
}

PAGE_CSS = f"""
@page {{ size: {W}px {H}px; margin: 0 }}
html, body {{ margin: 0; padding: 0; background: #1F2A33 }}
section {{ position: relative; width: {W}px; height: {H}px; box-sizing: border-box; overflow: hidden;
          page-break-after: always; break-after: page }}
section * {{ box-sizing: border-box }}
h1, h2, h3, p, ul, ol, table {{ margin: 0 }}
ul, ol {{ padding-left: 1.2em }}
table {{ border-collapse: collapse; width: 100% }}
th, td {{ padding: 0.35em 0.6em; border-bottom: 1px solid #D6D1C4 }}
th {{ font-weight: 600; text-align: left }}
aside {{ display: none }}
"""


def load_deck():
    with open(os.path.join(HERE, "deck.json")) as f:
        deck = json.load(f)
    slides = []
    for sid in deck["order"]:
        with open(os.path.join(HERE, "slides", f"{sid}.html")) as f:
            html = f.read()
        for blob, path in FIGURES.items():
            html = html.replace(f"/_blob/{blob}", "file://" + os.path.join(BASE_DIR, path))
        missing = re.findall(r"/_blob/([0-9a-f]+)", html)
        if missing:
            raise ValueError(f"slide {sid}: unknown figure ids {missing}")
        notes = re.search(r"<aside>(.*?)</aside>", html, re.S)
        slides.append({"id": sid, "html": html, "notes": notes.group(1).strip() if notes else ""})
    return deck, slides


def page(deck, sections):
    fonts = "".join(f'<link rel="stylesheet" href="{f["href"]}">' for f in deck["faces"].values() if "href" in f)
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{deck["title"]}</title>{fonts}'
            f"<style>{PAGE_CSS}</style></head><body>{''.join(sections)}</body></html>")


def chrome(*args):
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--allow-file-access-from-files",
                    "--virtual-time-budget=8000", *args], check=True, capture_output=True)


def main():
    deck, slides = load_deck()
    out_pdf = os.path.join(HERE, "fireWild_phase1.pdf")
    out_pptx = os.path.join(HERE, "fireWild_phase1.pptx")
    tmp = tempfile.mkdtemp()
    try:
        full = os.path.join(tmp, "deck.html")
        with open(full, "w") as f:
            f.write(page(deck, [s["html"] for s in slides]))
        chrome("--no-pdf-header-footer", f"--print-to-pdf={out_pdf}", "file://" + full)
        print(f"wrote {out_pdf}")

        from pptx import Presentation
        from pptx.util import Emu
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in (16:9)
        for s in slides:
            one = os.path.join(tmp, f"{s['id']}.html")
            png = os.path.join(tmp, f"{s['id']}.png")
            with open(one, "w") as f:
                f.write(page(deck, [s["html"]]))
            chrome(f"--window-size={W},{H}", f"--screenshot={png}", "file://" + one)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
            slide.notes_slide.notes_text_frame.text = s["notes"]
        prs.save(out_pptx)
        print(f"wrote {out_pptx} ({len(slides)} slides)")
    finally:
        shutil.rmtree(tmp)


if __name__ == "__main__":
    main()
